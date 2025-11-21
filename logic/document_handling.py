from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import os
import re


class documents:
    @staticmethod
    def save_scripts_to_docx(scripts, no_of_sequences, language):
        # Create 'data' folder if it doesn't exist
        data_dir = os.path.join(os.path.dirname(__file__), "data")
        os.makedirs(data_dir, exist_ok=True)

        # Determine the next available number for the current language
        prefix = f"script_{language}_"
        existing_files = [f for f in os.listdir(data_dir) if f.startswith(prefix) and f.endswith(".docx")]
        max_number = 0
        for fname in existing_files:
            match = re.match(rf"{re.escape(prefix)}(\d+)\.docx", fname)
            if match:
                number = int(match.group(1))
                if number > max_number:
                    max_number = number
        next_number = max_number + 1

        doc = Document()
        for i, script in enumerate(scripts):
            scenario_name = f"Scenario {i+1}" if no_of_sequences > 1 else "Scenario"
            # Scenario title, bold
            p = doc.add_paragraph()
            run = p.add_run(scenario_name)
            run.bold = True
            run.font.size = Pt(14)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

            # Extract summary and setting
            script_str = str(script)
            summary_pattern = r'^\[summary of script\]:(.*?)(?:\n|$)'
            settings_pattern = r'^\[settings\]:(.*?)(?:\n|$)'

            summary_match = re.search(summary_pattern, script_str, re.MULTILINE | re.IGNORECASE)
            settings_match = re.search(settings_pattern, script_str, re.MULTILINE | re.IGNORECASE)

            summary_text = summary_match.group(1).strip() if summary_match else ""
            setting_text = settings_match.group(1).strip() if settings_match else ""

            # Compose setting content = summary + optional setting
            content_lines = []
            if summary_text:
                content_lines.append(summary_text)
            if setting_text:
                content_lines.append(setting_text)
            content = "\n".join(content_lines)

            # "settings:" bold
            p_settings = doc.add_paragraph()
            run_settings = p_settings.add_run("settings: ")
            run_settings.bold = True
            run_settings.font.size = Pt(12)
            # Add extracted summary+settings
            p_settings.add_run(content)

            # horizontal line visual divider
            doc.add_paragraph().add_run("----------------------------------------")

            # Remove summary/settings lines from script content
            script_content = re.sub(summary_pattern, "", script_str, flags=re.MULTILINE | re.IGNORECASE)
            script_content = re.sub(settings_pattern, "", script_content, flags=re.MULTILINE | re.IGNORECASE)

            # script content; preserve line breaks
            for line in script_content.strip().splitlines():
                doc.add_paragraph(line)

            # Space between scenarios
            doc.add_paragraph("")

        docx_path = os.path.join(data_dir, f"{prefix}{next_number}.docx")
        doc.save(docx_path)